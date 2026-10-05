#!/usr/bin/env python3
"""Append-only run journal and last-run snapshot for NGAME-CM."""

from __future__ import annotations

import json
import os
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

from ngame_cm_config import load_cm_config, local_now, next_scheduled_local
from ngame_env import repo_root

RUNS_FILENAME = "ngame_runs.jsonl"
LAST_RUN_FILENAME = "ngame_last_run.json"
JOURNAL_DIR_ENV = "NGAME_JOURNAL_DIR"
TRAINING_MATRIX_FILE = "NGAME_Training_Matrix.xlsx"
TRAINING_TARGET_DAYS = 30
FRAUD_ANALYSIS_FILE = "NGAME_Fraud_Analysis.json"
MANAGEMENT_DASHBOARD_FILE = "management_dashboard.json"


def journal_dir() -> Path:
    override = os.environ.get(JOURNAL_DIR_ENV)
    if override:
        return Path(override)
    return Path(repo_root()) / "logs"


def runs_path() -> Path:
    return journal_dir() / RUNS_FILENAME


def last_run_path() -> Path:
    return journal_dir() / LAST_RUN_FILENAME


def read_last_run() -> Optional[Dict[str, Any]]:
    path = last_run_path()
    if not path.is_file():
        return None
    try:
        with path.open(encoding="utf-8") as fh:
            data = json.load(fh)
        return data if isinstance(data, dict) else None
    except Exception:
        return None


def snapshot_pipeline_facts(root: Optional[Path] = None) -> Dict[str, Any]:
    """Read current matrix day count and risk without changing the detector."""
    base = Path(root) if root is not None else Path(repo_root())
    days = _count_training_days(base / TRAINING_MATRIX_FILE)
    if not (base / TRAINING_MATRIX_FILE).exists():
        phase = "training"
    else:
        phase = "fraud_analysis" if days >= TRAINING_TARGET_DAYS else "training"
    return {
        "phase": phase,
        "matrix_days": days,
        "risk": _snapshot_risk(base),
    }


def write_run_record(
    *,
    trigger: str,
    ok: bool,
    exit_code: int,
    skipped_reason: Optional[str] = None,
    phase: Optional[str] = None,
    matrix_days: Optional[int] = None,
    risk: Optional[str] = None,
    extra: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    """
    Append one JSON object to logs/ngame_runs.jsonl and refresh last-run JSON.

    A skipped already-ran-today is success for uptime evidence, not a training
    day. A lock-busy from an overlapping dashboard click is not a missed schedule.
    """
    cfg = load_cm_config()
    now_utc = datetime.now(timezone.utc).replace(microsecond=0)
    now_local = local_now(cfg)
    facts = snapshot_pipeline_facts()
    record: Dict[str, Any] = {
        "utc_ts": now_utc.isoformat().replace("+00:00", "Z"),
        "local_date": now_local.date().isoformat(),
        "phase": phase if phase is not None else facts["phase"],
        "trigger": trigger,
        "ok": bool(ok),
        "exit_code": int(exit_code),
        "skipped_reason": skipped_reason,
        "matrix_days": matrix_days if matrix_days is not None else facts["matrix_days"],
        "risk": risk if risk is not None else facts["risk"],
    }
    if extra:
        record.update(extra)

    directory = journal_dir()
    directory.mkdir(parents=True, exist_ok=True)
    with runs_path().open("a", encoding="utf-8") as fh:
        fh.write(json.dumps(record, ensure_ascii=True) + "\n")

    last = dict(record)
    last["next_scheduled_local"] = next_scheduled_local(cfg)
    with last_run_path().open("w", encoding="utf-8") as fh:
        json.dump(last, fh, indent=2)
        fh.write("\n")
    return record


def _count_training_days(matrix_path: Path) -> int:
    """Day columns start at column 4 (A=type, B=μ, C=σ). Same rule as the dashboard."""
    if not matrix_path.exists():
        return 0
    try:
        from openpyxl import load_workbook

        wb = load_workbook(matrix_path, read_only=True)
        ws = wb.active
        days = max((ws.max_column or 0) - 3, 0)
        wb.close()
        return days
    except Exception:
        return 0


def _load_json(path: Path) -> Dict[str, Any]:
    try:
        with path.open(encoding="utf-8") as fh:
            data = json.load(fh)
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def _snapshot_risk(root: Path) -> Optional[str]:
    fraud = _load_json(root / FRAUD_ANALYSIS_FILE)
    management = _load_json(root / MANAGEMENT_DASHBOARD_FILE)
    management_summary = (
        management.get("management_summary")
        or (fraud.get("warning_result") or {}).get("management_summary")
        or {}
    )
    risk = (
        management_summary.get("overall_risk_level")
        or (fraud.get("summary") or {}).get("overall_risk_level")
        or (fraud.get("warning_result") or {}).get("overall_risk_level")
        or (fraud.get("risk_assessment") or {}).get("overall_risk")
    )
    if risk in (None, ""):
        return None
    return str(risk)


def read_run_records() -> List[Dict[str, Any]]:
    """All jsonl run records, skipping corrupt lines."""
    return _read_run_records()


def _read_run_records() -> List[Dict[str, Any]]:
    path = runs_path()
    if not path.is_file():
        return []
    records: List[Dict[str, Any]] = []
    try:
        with path.open(encoding="utf-8") as fh:
            for line in fh:
                line = line.strip()
                if not line:
                    continue
                try:
                    item = json.loads(line)
                except json.JSONDecodeError:
                    continue
                if isinstance(item, dict):
                    records.append(item)
    except OSError:
        return []
    return records


def consecutive_scheduled_success_days() -> int:
    """
    Trailing calendar days on which the scheduler succeeded (including
    already_ran / wrong_weekday). lock_busy does not break the streak.
    """
    return len(scheduled_success_streak_dates())


def scheduled_success_streak_dates() -> List[str]:
    cfg = load_cm_config()
    facts = snapshot_pipeline_facts()
    weekdays_only = facts["phase"] == "training" and bool(cfg.get("phase1_weekdays_only"))
    by_date: Dict[str, Dict[str, bool]] = {}
    for record in _read_run_records():
        if record.get("trigger") != "schedule":
            continue
        day = record.get("local_date")
        if not day:
            continue
        bucket = by_date.setdefault(day, {"success": False, "hard_fail": False})
        if record.get("ok"):
            bucket["success"] = True
        elif record.get("skipped_reason") != "lock_busy":
            bucket["hard_fail"] = True
    if not by_date:
        return []
    return _scheduled_success_streak_dates(by_date, weekdays_only)


def _scheduled_success_streak_dates(
    by_date: Dict[str, Dict[str, bool]],
    weekdays_only: bool,
) -> List[str]:
    from datetime import date, timedelta

    cursor = date.fromisoformat(max(by_date))
    dates: List[str] = []
    for _ in range(400):
        key = cursor.isoformat()
        info = by_date.get(key)
        if info and info["success"]:
            dates.append(key)
        elif info and info["hard_fail"]:
            break
        elif info:
            pass
        else:
            if weekdays_only and cursor.weekday() >= 5:
                pass
            else:
                break
        cursor = cursor - timedelta(days=1)
    dates.reverse()
    return dates


def lamp_last_result(last: Optional[Dict[str, Any]]) -> Optional[str]:
    if not last:
        return None
    if not last.get("ok"):
        return "FAIL"
    skip = last.get("skipped_reason")
    if skip in ("already_ran", "wrong_weekday"):
        return "OK"
    risk = last.get("risk")
    if risk:
        return str(risk).upper()
    return "OK"


def lamp_status() -> Dict[str, Any]:
    """FRP lamp: last run, result, streak, next scheduled — no click required."""
    cfg = load_cm_config()
    last = read_last_run()
    return {
        "last_run_utc": (last or {}).get("utc_ts"),
        "last_run_local_date": (last or {}).get("local_date"),
        "last_result": lamp_last_result(last),
        "last_ok": None if last is None else bool(last.get("ok")),
        "last_trigger": (last or {}).get("trigger"),
        "skipped_reason": (last or {}).get("skipped_reason"),
        "consecutive_scheduled_success_days": consecutive_scheduled_success_days(),
        "next_scheduled_local": (last or {}).get("next_scheduled_local")
        or next_scheduled_local(cfg),
        "scheduler_enabled": bool(cfg.get("enabled")),
        "dashboard_is_launcher": bool(cfg.get("dashboard_is_launcher")),
    }
